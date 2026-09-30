import uuid
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app import app, store
from rag.ingest import chunk_sections, extract_file

client = TestClient(app)

PASSWORD = "StrongPass1!"


def _unique_email() -> str:
    return f"user{uuid.uuid4().hex[:10]}@example.com"


def signup(client: TestClient, email: str) -> str:
    """Sign up and verify the account; returns the verified user's email."""
    r = client.post("/api/auth/signup", json={"name": "Test User", "email": email, "password": PASSWORD})
    assert r.status_code == 200, r.text
    code = r.json().get("dev_code")
    assert code, "Expected dev code (SMTP not configured)"
    vr = client.post("/api/auth/otp/verify", json={"email": email, "purpose": "signup", "code": code})
    assert vr.status_code == 200, vr.text
    assert vr.json()["redirect"] == "/login?verified=1"
    return email


def login(client: TestClient, email: str) -> TestClient:
    r = client.post("/api/auth/login", json={"email": email, "password": PASSWORD})
    assert r.status_code == 200, r.text
    return client


def auth_client() -> TestClient:
    """A TestClient logged in with a fresh verified account."""
    c = TestClient(app)
    email = signup(c, _unique_email())
    return login(c, email)


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------
def test_signup_login_and_me():
    c = TestClient(app)
    email = signup(c, _unique_email())
    login(c, email)
    me = c.get("/api/auth/me")
    assert me.status_code == 200
    assert me.json()["user"]["email"] == email


def test_duplicate_signup_rejected():
    c = TestClient(app)
    email = _unique_email()
    signup(c, email)
    r = c.post("/api/auth/signup", json={"name": "Other", "email": email, "password": PASSWORD})
    assert r.status_code == 409


def test_signup_does_not_expose_otp_when_delivery_fails(monkeypatch):
    monkeypatch.setenv("ALLOW_DEV_OTP", "false")
    monkeypatch.setattr("app.send_otp_email", lambda *args, **kwargs: False)
    r = TestClient(app).post(
        "/api/auth/signup",
        json={"name": "Mail Fail", "email": _unique_email(), "password": PASSWORD},
    )
    assert r.status_code == 503
    assert "dev_code" not in r.text


def test_login_wrong_password_rejected():
    c = TestClient(app)
    email = signup(c, _unique_email())
    r = c.post("/api/auth/login", json={"email": email, "password": "wrong-pass-1"})
    assert r.status_code == 401


def test_otp_three_wrong_attempts_lock():
    c = TestClient(app)
    email = _unique_email()
    r = c.post("/api/auth/signup", json={"name": "Lock Test", "email": email, "password": PASSWORD})
    code = r.json()["dev_code"]
    # First two wrong attempts
    for _ in range(2):
        res = c.post("/api/auth/otp/verify", json={"email": email, "purpose": "signup", "code": "000000"})
        assert res.status_code == 400
        assert "attempt(s) remaining" in res.json()["detail"]
    # Third wrong attempt invalidates the code and requires a new one.
    res = c.post("/api/auth/otp/verify", json={"email": email, "purpose": "signup", "code": "000000"})
    assert res.status_code == 429
    assert "new verification code" in res.json()["detail"]
    # The old correct code is rejected after invalidation.
    locked = c.post("/api/auth/otp/verify", json={"email": email, "purpose": "signup", "code": code})
    assert locked.status_code == 400


def test_api_requires_session():
    assert client.get("/api/documents").status_code == 401
    assert client.post("/api/ask", json={"question": "hi"}).status_code == 401
    assert client.post("/api/ingest/url", json={"url": "https://example.com"}).status_code == 401


def test_api_key_grants_headless_access():
    c = TestClient(app)
    c.headers.update({"X-API-Key": "test-api-key"})
    assert c.get("/api/documents").status_code == 200
    bad = TestClient(app)
    bad.headers.update({"X-API-Key": "wrong-key"})
    assert bad.get("/api/documents").status_code == 401


def test_logout_invalidates_session():
    c = auth_client()
    assert c.get("/api/auth/me").status_code == 200
    assert c.post("/api/auth/logout").status_code == 200
    assert c.get("/api/auth/me").status_code == 401


# ---------------------------------------------------------------------------
# Extraction (unchanged behaviour)
# ---------------------------------------------------------------------------
def test_text_markdown_and_csv_extraction():
    md = extract_file("guide.md", b"# Dental Guide\nRoot canal treatment saves an infected tooth.")
    csv = extract_file("doctors.csv", b"name,speciality\nSara,Orthodontics\n")
    assert "Root canal" in md[0].text
    assert "Sara" in csv[0].text


def test_docx_and_pptx_extraction(tmp_path: Path):
    doc_path = tmp_path / "sample.docx"
    doc = Document(); doc.add_paragraph("Clinic opens at nine."); doc.save(doc_path)
    ppt_path = tmp_path / "sample.pptx"
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = "Appointments"; slide.placeholders[1].text = "Thirty minute slots"; prs.save(ppt_path)
    assert "opens at nine" in extract_file(doc_path.name, doc_path.read_bytes())[0].text
    assert "Thirty minute" in extract_file(ppt_path.name, ppt_path.read_bytes())[0].text


def test_chunk_overlap_does_not_loop():
    chunks = chunk_sections("long.txt", extract_file("long.txt", ("knowledge " * 500).encode()), size=200, overlap=30)
    assert 10 < len(chunks) < 40


# ---------------------------------------------------------------------------
# RAG API (authenticated)
# ---------------------------------------------------------------------------
def test_api_ingest_retrieve_and_delete(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    c = auth_client()
    source = "test-knowledge.md"
    response = c.post("/api/ingest/files", files=[("files", (source, b"The clinic is closed every Sunday.", "text/markdown"))])
    assert response.status_code == 200
    assert response.json()["results"][0]["status"] == "queued"
    answer = c.post("/api/ask", json={"question": "When is the clinic closed?", "top_k": 3})
    assert answer.status_code == 200
    assert any(s["source"] == source for s in answer.json()["sources"])
    deleted = c.post("/api/documents/delete", json={"source": source})
    assert deleted.status_code == 200


def test_selected_sources_exclude_unselected_documents(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    c = auth_client()
    selected = "selected-policy.md"
    excluded = "excluded-policy.md"
    files = [
        ("files", (selected, b"Project Falcon uses the blue approval procedure.", "text/markdown")),
        ("files", (excluded, b"Project Falcon uses the red confidential procedure.", "text/markdown")),
    ]
    assert c.post("/api/ingest/files", files=files).status_code == 200
    response = c.post(
        "/api/ask",
        json={"question": "Which procedure does Project Falcon use?", "sources": [selected]},
    )
    data = response.json()
    assert response.status_code == 200
    assert data["selected_sources"] == [selected]
    assert data["sources"]
    assert all(item["source"] == selected for item in data["sources"])
    assert excluded not in data["answer"]
    c.post("/api/documents/delete", json={"source": selected})
    c.post("/api/documents/delete", json={"source": excluded})


def test_health_and_home():
    assert client.get("/health").json()["status"] == "ok"
    assert "Synapse" in client.get("/").text
    assert client.get("/login").status_code == 200
    assert client.get("/signup").status_code == 200
    assert client.get("/verify").status_code == 200


def test_ingest_folder_scans_recursively(tmp_path):
    (tmp_path / "docs").mkdir()
    (tmp_path / "docs" / "readme.md").write_text("Folder scanner stores this paragraph about blueprints.", encoding="utf-8")
    (tmp_path / "notes.txt").write_text("Recursive walk also catches files in the root folder.", encoding="utf-8")
    (tmp_path / "ignore.bin").write_bytes(b"\x00\x01binary")
    c = auth_client()
    r = c.post("/api/ingest/folder", json={"path": str(tmp_path)})
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["files_found"] == 2
    assert data["chunk_count"] > 0
    assert data["stored_chunks"] == data["chunk_count"]
    sources = {item["source"] for item in data["indexed"]}
    assert "docs\\readme.md" in sources or "docs/readme.md" in sources
    assert "notes.txt" in sources
    assert data["errors"] == []
    for item in data["indexed"]:
        c.post("/api/documents/delete", json={"source": item["source"]})


def test_ingest_folder_rejects_missing_path():
    c = auth_client()
    r = c.post("/api/ingest/folder", json={"path": "Z:/does/not/exist"})
    assert r.status_code == 400
